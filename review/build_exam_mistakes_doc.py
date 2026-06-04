from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "review/考前易错点总结_C++上机.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for i, width in enumerate(widths):
            set_cell_width(row.cells[i], width)
            set_cell_margins(row.cells[i])
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, east_asia="Microsoft YaHei", latin="Calibri", size=None, bold=None, color=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para_font(paragraph, size=None, bold=None, color=None):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    size = {1: 16, 2: 13, 3: 12}.get(level, 12)
    color = "2E74B5" if level in (1, 2) else "1F4D78"
    set_run_font(run, size=size, bold=True, color=color)
    return p


def add_body(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=10.5)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run, size=10.5)


def add_matrix_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "E8EEF5")
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_para_font(p, size=10, bold=True, color="1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                set_para_font(p, size=9.5)
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.2
for name, size, color in [
    ("Heading 1", 16, "2E74B5"),
    ("Heading 2", 13, "2E74B5"),
    ("Heading 3", 12, "1F4D78"),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 16)
    st.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("C++ 上机考前易错点总结")
set_run_font(r, size=22, bold=True, color="0B2545")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("基于 chapter 7-11、simulate 2、Chapter 8 复盘与 Date 类上机题整理")
set_run_font(r, size=10.5, color="555555")

add_heading(doc, "一、最优先盯住的 12 个易错点", 1)
priority_rows = [
    ("1", "看 main 倒推接口", "先圈表达式：cout<<obj、cin>>obj、obj[i]、a+b、++obj、Base* p=new Derived。接口签名错了，逻辑写对也没用。"),
    ("2", "+ / - 不要改左操作数", "除非题目明确要求，operator+ 和 operator- 应复制一个临时对象返回；Date、Time 这类题最容易把原对象改坏。"),
    ("3", "把规则抽成小函数", "日期、时间、分数、矩阵、金额不要堆复杂 if/switch。先写 isLeap、daysInMonth、normalize、gcd、index。"),
    ("4", "switch 必须 break", "大月小月、菜单分支、类型分支都容易因为忘 break 穿透。能用 daysInMonth 函数时少用大段 switch。"),
    ("5", "格式输出是硬分", "setw/setfill、fixed/setprecision、boolalpha、括号逗号空格、YYYY-MM-DD 都要完全按样例。"),
    ("6", "模板里别写死 int", "VectorSum、Matrix、ProductVector 里临时变量应是 T；参数优先 const vector<T>&。"),
    ("7", "比较函数相等返回 false", "sort 的 cmp 表达“a 是否在 b 前面”。优先级/耗时都相同时必须 false，避免排序规则不严格。"),
    ("8", "vector 下标和边界", "pop_back 前先确认非空；erase 用 begin()+index；size_t 与 int 混用、空串倒序遍历都要小心。"),
    ("9", "二维 vector 初始化", "vector<vector<T>>(row, vector<T>(col))；矩阵加法维度用左矩阵行列，乘法临时变量用 T。"),
    ("10", "继承/抽象类签名完全一致", "纯虚函数和派生类函数的返回值、名字、参数必须完全一样；必要时加 override 让编译器帮你查。"),
    ("11", "金额/分数不用浮点硬算", "RMB 用总分 cents，Fraction 用 gcd 约分并保证分母为正；double 只用于最终显示或题目明确要求。"),
    ("12", "输入方式别混乱", "一整行用 getline + stringstream；cin 后接 getline 时处理换行；用到 toupper、sort、stringstream 要记头文件。"),
]
add_matrix_table(doc, ["优先级", "易错点", "复习提醒"], priority_rows, [900, 2100, 6360])

add_heading(doc, "二、按题型归纳", 1)

add_heading(doc, "1. 运算符重载", 2)
add_body(doc, "你的 Chapter 8 复盘显示，语法基本能写出来，真正危险的是“运算符语义不像内置类型”。")
add_bullets(doc, [
    "算术运算：返回新对象，通常不改变两个操作数。Date::operator+、Time::operator- 这类题要先复制临时对象。",
    "流输出：friend ostream& operator<<(ostream& out, const T& obj)，内部只用 out，不要混用 cout。",
    "流输入：friend istream& operator>>(istream& in, T& obj)，第二个参数必须是非 const 引用。",
    "下标运算：如果 m[i][j] 或 obj[i]=x 要成立，operator[] 要返回引用。",
    "比较运算：只读函数加 const，参数用 const 引用；不要漏 return。",
])
add_body(doc, "固定签名：")
code = [
    "ClassName operator+(const ClassName& rhs) const;",
    "bool operator<(const ClassName& rhs) const;",
    "friend ostream& operator<<(ostream& out, const ClassName& obj);",
    "friend istream& operator>>(istream& in, ClassName& obj);",
    "RowType& operator[](int i);",
]
for line in code:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(line)
    set_run_font(r, latin="Consolas", east_asia="Microsoft YaHei", size=9.5)

add_heading(doc, "2. 日期、时间、进位类题", 2)
add_body(doc, "你回忆的 Date 类题思路本质上是对的：每天推进并跨月跨年。但把大月、小月、二月拆成很多分支，会把错误点放大。最稳做法是把“这个月有多少天”抽出来。")
add_bullets(doc, [
    "闰年公式：year%400==0 或 year%4==0 且 year%100!=0。",
    "daysInMonth(y,m) 返回 28/29/30/31，主逻辑只问这个函数。",
    "++ 只负责加一天；+k 复制对象后循环调用 ++。",
    "switch 写月份时每个 case 后要 break；多个月份判断必须用 ||，不能写 month==1||3。",
    "YYYY-MM-DD 用 setw(4)、setw(2)、setfill('0')，不要输出空格版日期。",
])

add_heading(doc, "3. 模板类 / 函数模板", 2)
add_bullets(doc, [
    "题目给的是函数模板形式时，尽量照抄：vector<T> f(const vector<T>& input)。",
    "模板内部不要把 result、temp、max 写成 int；应写 T result = 0 / 1。",
    "字符串模板用 basic_string<T>，不要明明 template<typename T> 却固定 string。",
    "模板派生类访问基类成员时，用 this->elements 或 using Base<T>::elements。",
    "非类型模板参数如 template<typename T, size_t MAX_SIZE>，题目强调固定数组时最好真的体现 MAX_SIZE。",
])

add_heading(doc, "4. 继承、多态、抽象类", 2)
add_bullets(doc, [
    "纯虚函数名、返回值、参数必须和派生类实现完全一致；建议写 override。",
    "基类指针调用派生类函数时，基类函数必须 virtual。",
    "有 new 就考虑 delete；基类指针 delete 派生对象时，规范上基类析构函数应 virtual。",
    "多继承/虚继承构造时，最底层派生类负责初始化虚基类。",
    "构造函数初始化列表按继承链写，间接基类一般由直接基类负责初始化。",
])

add_heading(doc, "5. STL、字符串、数组边界", 2)
add_bullets(doc, [
    "手动删除 vector 元素时，不要自己整体左移再忘 pop_back；优先 items.erase(items.begin()+index)。",
    "pop_back、back、arr[0] 前确认容器非空。",
    "倒序遍历 string/vector 时，size() 是无符号类型，空字符串会出问题。",
    "路径简化、身份证补位、大数运算这类题，最容易错在 i+1、i+2 越界和结尾处理。",
    "字符数组需要 '\\0' 结尾；不要用 NULL 判断 char 内容。",
])

add_heading(doc, "6. 数值与数学细节", 2)
add_bullets(doc, [
    "几何题 double 不要直接 ==0，用 EPS，如 fabs(x)<1e-9。",
    "pow 返回 double，整数进制/多项式/大数计算尽量用整数累乘。",
    "分数约分用 gcd(abs(num), den)，分母始终为正。",
    "金额类用 long long cents，不要用 double 当核心状态。",
    "矩阵乘法的累加变量跟着模板类型走：T temp = 0。",
])

add_heading(doc, "三、你最容易重复犯的个人模式", 1)
pattern_rows = [
    ("把分支写复杂", "Date 题用大月/小月/闰年 switch + 多个 while，容易忘 break。", "先抽 daysInMonth，再统一推进。"),
    ("题意没有先圈接口", "函数参数是指针、引用、const、模板，容易写成自己顺手的形式。", "先照 main 写函数声明，再补实现。"),
    ("只过样例，没想隐藏数据", "Time 只处理一次进位、Triangle 样例没暴露向量写错、Ternary 0 空输出。", "提交前自己造极端数据：0、负数、跨月、相等、空集合。"),
    ("返回值语义不自然", "operator+ 或 operator- 修改了当前对象。", "凡是 +、- 先问：内置 int 的 a+b 会不会改 a？"),
    ("模板题仍按 int 题写", "result、threshold、temp 写死 int。", "看到 template，变量类型先优先 T。"),
    ("格式靠感觉", "括号、逗号、补零、固定两位小数、boolalpha 容易遗漏。", "复制样例输出格式，逐字符核对。"),
]
add_matrix_table(doc, ["模式", "表现", "应对动作"], pattern_rows, [1700, 3600, 4060])

add_heading(doc, "四、上机最后 5 分钟检查清单", 1)
add_numbered(doc, [
    "main 中每一个表达式都能匹配你的函数签名吗？尤其是 const、引用、指针、friend、static。",
    "所有非 void 函数是否每条路径都有 return？比较函数相等时是否返回 false？",
    "+、-、* 是否该返回新对象？有没有不小心改 this？",
    "输出格式是否和样例完全一致：空格、换行、括号、冒号、小数位、补零。",
    "是否包含必要头文件：algorithm、sstream、cctype、iomanip、cmath。",
    "vector/string 是否可能空？pop_back、back、[0]、倒序循环是否安全？",
    "模板代码里的临时变量是否写死 int？参数是否可以改成 const 引用？",
    "日期/时间/分数/金额有没有统一的辅助函数或内部表示？",
    "new 出来的对象是否 delete？数组 new[] 是否 delete[]？",
    "自己额外测一组边界：跨月、跨年、2 月 28/29、相等、负数、空结果、重复元素。",
])

add_heading(doc, "五、可直接背的安全模板", 1)

add_heading(doc, "Date 类核心模板", 2)
date_template = [
    "bool isLeap(int y) { return y%400==0 || (y%4==0 && y%100!=0); }",
    "int daysInMonth(int y,int m) {",
    "    if (m==2) return isLeap(y)?29:28;",
    "    if (m==4 || m==6 || m==9 || m==11) return 30;",
    "    return 31;",
    "}",
    "Date& operator++() {",
    "    d++;",
    "    if (d > daysInMonth(y,m)) { d=1; m++; if (m>12) { m=1; y++; } }",
    "    return *this;",
    "}",
    "Date operator+(int k) const { Date res=*this; while(k--) ++res; return res; }",
]
for line in date_template:
    p = doc.add_paragraph()
    r = p.add_run(line)
    set_run_font(r, latin="Consolas", east_asia="Microsoft YaHei", size=9)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(1)

add_heading(doc, "sort 比较函数模板", 2)
cmp_template = [
    "bool cmp(const Task* a, const Task* b) {",
    "    if (a->priority != b->priority) return a->priority > b->priority;",
    "    if (a->getTime() != b->getTime()) return a->getTime() < b->getTime();",
    "    return false;",
    "}",
]
for line in cmp_template:
    p = doc.add_paragraph()
    r = p.add_run(line)
    set_run_font(r, latin="Consolas", east_asia="Microsoft YaHei", size=9)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(1)

add_heading(doc, "六、来源范围", 1)
add_bullets(doc, [
    "review/simulate2_review.md：12 道模拟题复盘，重点是模板继承、集合、矩阵、库存、格式和边界。",
    "review/chapter11_review.md：10 道模板题复盘，重点是 const 引用、T 类型、basic_string、固定数组和二分查找。",
    "review/Chapter8_运算符重载_题目与代码复盘.docx：7 道运算符重载题复盘，重点是无副作用、const、引用、内部表示。",
    "chapter 7-10 源码：类与对象、静态成员、继承多态、抽象类、矩阵/日期/金额/几何/字符串处理等代码。",
    "上周 Date 类上机题：重点提醒不要把日期推进写成复杂 switch 分支，改用 daysInMonth + ++。",
])

doc.core_properties.title = "C++ 上机考前易错点总结"
doc.core_properties.subject = "根据个人代码与复盘整理"
doc.core_properties.author = "Codex"
doc.save(OUT)
print(OUT)
